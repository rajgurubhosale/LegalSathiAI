import re
from pathlib import Path
from uuid import uuid4

import docx
import fitz  # PyMuPDF
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

from src.utils.main_utils import read_config_file
from src.logger import logger


def clean_text(text: str) -> str:
    """Cleans raw extracted text: removes excessive spaces and blank lines."""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def parse_document(file_path: Path) -> list[tuple[int, str]]:
    """
    Parses PDF or DOCX file and returns a list of (page_number, text).
    """
    ext = file_path.suffix.lower()
    pages: list[tuple[int, str]] = []

    # 1. Parse PDF using PyMuPDF (fitz)
    if ext == ".pdf":
        doc = fitz.open(file_path)
        for page_num in range(len(doc)):
            raw_text = doc[page_num].get_text()
            cleaned = clean_text(raw_text)
            if cleaned:
                pages.append((page_num + 1, cleaned))
        doc.close()

    # 2. Parse DOCX using python-docx
    elif ext == ".docx":
        doc = docx.Document(file_path)
        full_text = "\n\n".join(
            p.text.strip() for p in doc.paragraphs if p.text.strip()
        )
        cleaned = clean_text(full_text)
        if cleaned:
            pages.append((1, cleaned))

    return pages


class UserVault:
    _shared_embedding_model = None

    def __init__(self, embedding_model=None):
        self.config = read_config_file()

        vault_cfg = self.config.get("user_vault", {})
        collection_name = vault_cfg.get("collection_name", "users_personal_vault")
        db_path = vault_cfg.get("db_path", "data/user_vault_vectorstore")
        self.chunk_size = vault_cfg.get("chunk_size", 1000)
        self.chunk_overlap = vault_cfg.get("chunk_overlap", 100)

        # 1. Reuse existing embedding model if passed
        if embedding_model is not None:
            UserVault._shared_embedding_model = embedding_model

        # 2. Only load if not already in cache
        if UserVault._shared_embedding_model is None:
            model_name = self.config["embedding"]["model_name"]
            UserVault._shared_embedding_model = HuggingFaceEmbeddings(
                model_name=model_name,
                encode_kwargs={"normalize_embeddings": True},
            )

        self.embedding_model = UserVault._shared_embedding_model
        self.store = Chroma(
            collection_name=collection_name,
            embedding_function=self.embedding_model,
            persist_directory=str(db_path),
            create_collection_if_not_exists=True,
        )

        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
        )

    def add_documents(self, documents: list[Document], batch_size: int = 64):
        """Adds LangChain documents to Chroma in safe batches."""
        for i in range(0, len(documents), batch_size):
            batch = documents[i : i + batch_size]
            ids = [doc.metadata["chunk_id"] for doc in batch]
            self.store.add_documents(batch, ids=ids)
            logger.info(f"Stored {min(i + batch_size, len(documents))}/{len(documents)} chunks in vault")

    def process_and_store(self, file_path: Path | str, user_id: int, doc_name: str) -> dict:
        """
        End-to-end ingestion:
        1. Parses PDF (PyMuPDF) or DOCX (python-docx)
        2. Cleans text
        3. Chunks into 1000-100 characters recursively
        4. Attaches metadata: user_id, doc_name, page_no, chunk_id
        5. Embeds and stores in ChromaDB
        """
        path = Path(file_path)
        if not path.exists():
            return {"status": "error", "message": f"File not found: {file_path}"}

        # 1. Parse & Clean
        pages = parse_document(path)
        if not pages:
            return {"status": "empty", "message": "No readable text extracted", "chunks_added": 0}

        # 2. Chunk & Enrich with metadata
        documents: list[Document] = []
        batch_id = uuid4().hex[:8]

        for page_no, page_text in pages:
            chunks = self.text_splitter.split_text(page_text)
            for idx, chunk in enumerate(chunks):
                documents.append(
                    Document(
                        page_content=chunk,
                        metadata={
                            "user_id": user_id,
                            "doc_name": doc_name,
                            "page_no": page_no,
                            "chunk_id": f"{user_id}_{batch_id}_{page_no}_{idx}",
                        },
                    )
                )

        # 3. Embed & Store in Chroma
        if documents:
            self.add_documents(documents)

        return {
            "status": "done",
            "doc_name": doc_name,
            "user_id": user_id,
            "chunks_added": len(documents),
        }

    def process_files(self, files: list[tuple[Path | str, str]], user_id: int) -> int:
        """
        Parses multiple files, gathers all chunks, and embeds them into Chroma in safe batches.
        files: list of (file_path, doc_name) tuples
        """
        all_documents: list[Document] = []
        batch_id = uuid4().hex[:8]

        for file_path, doc_name in files:
            path = Path(file_path)
            if not path.exists():
                continue
            pages = parse_document(path)
            for page_no, page_text in pages:
                chunks = self.text_splitter.split_text(page_text)
                for idx, chunk in enumerate(chunks):
                    all_documents.append(
                        Document(
                            page_content=chunk,
                            metadata={
                                "user_id": user_id,
                                "doc_name": doc_name,
                                "page_no": page_no,
                                "chunk_id": f"{user_id}_{batch_id}_{page_no}_{idx}",
                            },
                        )
                    )

        if all_documents:
            self.add_documents(all_documents, batch_size=64)

        return len(all_documents)

    def search(self, query: str, user_id: int, k: int = 8):
        """Searches Chroma filtered strictly by user_id."""
        if self.store._collection.count() == 0:
            return []
            
        return self.store.similarity_search(
            query=query,
            k=k,
            filter={"user_id": user_id},
        )

    def delete_document(self, user_id: int, doc_name: str) -> int:
        """Deletes all chunks belonging to a specific document for a user from ChromaDB."""
        if self.store._collection.count() == 0:
            return 0
        try:
            matching = self.store._collection.get(where={"$and": [{"user_id": user_id}, {"doc_name": doc_name}]})
            ids = matching.get("ids", [])
            if ids:
                self.store._collection.delete(ids=ids)
            return len(ids)
        except Exception as e:
            logger.error(f"Error deleting chunks from vault: {e}")
            return 0

    def clear_vault(self, user_id: int) -> int:
        """Deletes all chunks belonging to a user from ChromaDB."""
        if self.store._collection.count() == 0:
            return 0
        try:
            matching = self.store._collection.get(where={"user_id": user_id})
            ids = matching.get("ids", [])
            if ids:
                self.store._collection.delete(ids=ids)
            return len(ids)
        except Exception as e:
            logger.error(f"Error clearing user vault from Chroma: {e}")
            return 0